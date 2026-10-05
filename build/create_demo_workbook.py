"""Create a fictional, complete workbook used by the browser walkthrough."""

from __future__ import annotations

import json
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "dist" / "Sistema_Financiero_Empresa_2026.xlsx"
OUTPUT_DIR = ROOT / "app" / "static" / "demo"
OUTPUT = OUTPUT_DIR / "Sistema_Financiero_Demo_2026.xlsx"
MANIFEST = OUTPUT_DIR / "manifest.json"

COMPANY = {
    "NombreEmpresa": "Comercial Andina Demo S.A.S.",
    "NIT": "900000000-1 (FICTICIO)",
    "RegimenTributario": "DEMO — validar y configurar para cada empresa",
    "ActividadEconomica": "Comercio y servicios (ejemplo)",
    "ResponsableIVA": "Configurar",
}

THIRD_PARTIES = [
    {
        "ID": "TER-DEMO-001", "Tipo": "Cliente", "NombreRazonSocial": "Cliente Aurora Demo S.A.S.",
        "NIT_ID": "900000001-2 (FICTICIO)", "Contacto": "Contacto de demostración",
        "Email": "cliente.demo@example.invalid", "Telefono": "0000000000",
        "Direccion": "Dirección ficticia", "Ciudad": "Bogotá D.C.", "Estado": "Activo",
        "Observaciones": "Registro ficticio para simulación; no usar como tercero real.",
    },
    {
        "ID": "TER-DEMO-002", "Tipo": "Proveedor", "NombreRazonSocial": "Suministros Pacífico Demo S.A.S.",
        "NIT_ID": "900000002-3 (FICTICIO)", "Contacto": "Contacto de demostración",
        "Email": "proveedor.demo@example.invalid", "Telefono": "0000000000",
        "Direccion": "Dirección ficticia", "Ciudad": "Medellín", "Estado": "Activo",
        "Observaciones": "Registro ficticio para simulación; no usar como tercero real.",
    },
    {
        "ID": "TER-DEMO-003", "Tipo": "Proveedor", "NombreRazonSocial": "Nube Clara Demo S.A.S.",
        "NIT_ID": "900000003-4 (FICTICIO)", "Contacto": "Contacto de demostración",
        "Email": "nube.demo@example.invalid", "Telefono": "0000000000",
        "Direccion": "Dirección ficticia", "Ciudad": "Cali", "Estado": "Activo",
        "Observaciones": "Registro ficticio para simulación; no usar como tercero real.",
    },
]

MOVEMENTS = [
    {
        "ID": "MOV-DEMO-001", "Fecha": date(2026, 9, 8), "Tipo": "Ingreso", "Documento": "FV-DEMO-001",
        "Tercero": "Cliente Aurora Demo S.A.S.", "NIT_ID": "900000001-2 (FICTICIO)",
        "Concepto": "Servicio mensual de demostración", "Categoria": "Ventas", "Subcategoria": "Servicios",
        "MedioPago": "Transferencia", "Cuenta": "Banco Principal", "Subtotal": 4_800_000,
        "Impuestos": 0, "Retenciones": 0, "Estado": "Pagado", "SoporteID": "DOC-DEMO-001",
        "Observaciones": "Operación completamente ficticia; valores sin cálculo tributario.",
    },
    {
        "ID": "MOV-DEMO-002", "Fecha": date(2026, 9, 12), "Tipo": "Venta", "Documento": "FV-DEMO-002",
        "Tercero": "Cliente Aurora Demo S.A.S.", "NIT_ID": "900000001-2 (FICTICIO)",
        "Concepto": "Venta de productos de demostración", "Categoria": "Productos", "Subcategoria": "General",
        "MedioPago": "Transferencia", "Cuenta": "Cuentas por Cobrar", "Subtotal": 2_150_000,
        "Impuestos": 0, "Retenciones": 0, "Estado": "Pendiente", "SoporteID": "DOC-DEMO-002",
        "Observaciones": "Saldo por cobrar de ejemplo.",
    },
    {
        "ID": "MOV-DEMO-003", "Fecha": date(2026, 9, 10), "Tipo": "Compra", "Documento": "FC-DEMO-003",
        "Tercero": "Suministros Pacífico Demo S.A.S.", "NIT_ID": "900000002-3 (FICTICIO)",
        "Concepto": "Suministros de oficina de demostración", "Categoria": "Inventario", "Subcategoria": "Mercancia",
        "MedioPago": "Transferencia", "Cuenta": "Banco Principal", "Subtotal": 850_000,
        "Impuestos": 0, "Retenciones": 0, "Estado": "Pagado", "SoporteID": "DOC-DEMO-003",
        "Observaciones": "Operación completamente ficticia; valores sin cálculo tributario.",
    },
    {
        "ID": "MOV-DEMO-004", "Fecha": date(2026, 9, 18), "Tipo": "Compra", "Documento": "FC-DEMO-004",
        "Tercero": "Suministros Pacífico Demo S.A.S.", "NIT_ID": "900000002-3 (FICTICIO)",
        "Concepto": "Equipo tecnológico de demostración", "Categoria": "Activos", "Subcategoria": "Tecnologia",
        "MedioPago": "Transferencia", "Cuenta": "Banco Principal", "Subtotal": 420_000,
        "Impuestos": 0, "Retenciones": 0, "Estado": "Pagado", "SoporteID": "DOC-DEMO-004",
        "Observaciones": "Operación completamente ficticia; valores sin cálculo tributario.",
    },
    {
        "ID": "MOV-DEMO-005", "Fecha": date(2026, 9, 21), "Tipo": "Gasto", "Documento": "FE-DEMO-005",
        "Tercero": "Nube Clara Demo S.A.S.", "NIT_ID": "900000003-4 (FICTICIO)",
        "Concepto": "Suscripción de software de demostración", "Categoria": "Administracion", "Subcategoria": "Software",
        "MedioPago": "Transferencia", "Cuenta": "Cuentas por Pagar", "Subtotal": 620_000,
        "Impuestos": 0, "Retenciones": 0, "Estado": "Pendiente", "SoporteID": "DOC-DEMO-005",
        "Observaciones": "Saldo por pagar de ejemplo.",
    },
]

DOCUMENTS = [
    {
        "ID": "DOC-DEMO-001", "Fecha": date(2026, 9, 8), "TipoDocumento": "Factura electronica", "Numero": "FV-DEMO-001",
        "Tercero": "Cliente Aurora Demo S.A.S.", "Identificacion": "900000001-2 (FICTICIO)",
        "Subtotal": 4_800_000, "Impuestos": 0, "Retenciones": 0, "Total": 4_800_000,
        "CUFE_Referencia": "CUFE-FICTICIO-DEMO-001", "Estado": "Validado",
        "RutaArchivo": "DEMO://documentos/FACTURA_DEMO_001.xml", "NombreArchivo": "FACTURA_DEMO_001.xml",
        "TipoArchivo": "XML", "FechaExtraccion": datetime(2026, 9, 8, 9, 15), "MetodoExtraccion": "XML",
        "Confianza": 1, "Observaciones": "XML estructurado ficticio; campos y sumas de ejemplo.",
    },
    {
        "ID": "DOC-DEMO-002", "Fecha": date(2026, 9, 12), "TipoDocumento": "Factura electronica", "Numero": "FV-DEMO-002",
        "Tercero": "Cliente Aurora Demo S.A.S.", "Identificacion": "900000001-2 (FICTICIO)",
        "Subtotal": 2_150_000, "Impuestos": 0, "Retenciones": 0, "Total": 2_150_000,
        "CUFE_Referencia": "CUFE-FICTICIO-DEMO-002", "Estado": "Validado",
        "RutaArchivo": "DEMO://documentos/FACTURA_DEMO_002.xml", "NombreArchivo": "FACTURA_DEMO_002.xml",
        "TipoArchivo": "XML", "FechaExtraccion": datetime(2026, 9, 12, 10, 5), "MetodoExtraccion": "XML",
        "Confianza": 1, "Observaciones": "Factura ficticia relacionada con una cuenta pendiente de cobro.",
    },
    {
        "ID": "DOC-DEMO-003", "Fecha": date(2026, 9, 10), "TipoDocumento": "Factura electronica", "Numero": "FC-DEMO-003",
        "Tercero": "Suministros Pacífico Demo S.A.S.", "Identificacion": "900000002-3 (FICTICIO)",
        "Subtotal": 850_000, "Impuestos": 0, "Retenciones": 0, "Total": 850_000,
        "CUFE_Referencia": "CUFE-FICTICIO-DEMO-003", "Estado": "Validado",
        "RutaArchivo": "DEMO://documentos/FACTURA_DEMO_003.xml", "NombreArchivo": "FACTURA_DEMO_003.xml",
        "TipoArchivo": "XML", "FechaExtraccion": datetime(2026, 9, 10, 11, 30), "MetodoExtraccion": "XML",
        "Confianza": 1, "Observaciones": "Factura ficticia estructurada; validación aritmética de ejemplo.",
    },
    {
        "ID": "DOC-DEMO-004", "Fecha": date(2026, 9, 18), "TipoDocumento": "Factura electronica", "Numero": "FC-DEMO-004",
        "Tercero": "Suministros Pacífico Demo S.A.S.", "Identificacion": "900000002-3 (FICTICIO)",
        "Subtotal": 420_000, "Impuestos": 0, "Retenciones": 0, "Total": 420_000,
        "CUFE_Referencia": "CUFE-FICTICIO-DEMO-004", "Estado": "Validado",
        "RutaArchivo": "DEMO://documentos/FACTURA_DEMO_004.xml", "NombreArchivo": "FACTURA_DEMO_004.xml",
        "TipoArchivo": "XML", "FechaExtraccion": datetime(2026, 9, 18, 14, 10), "MetodoExtraccion": "XML",
        "Confianza": 1, "Observaciones": "Factura ficticia de compra de activo.",
    },
    {
        "ID": "DOC-DEMO-005", "Fecha": date(2026, 9, 21), "TipoDocumento": "Factura electronica", "Numero": "FE-DEMO-005",
        "Tercero": "Nube Clara Demo S.A.S.", "Identificacion": "900000003-4 (FICTICIO)",
        "Subtotal": 620_000, "Impuestos": 0, "Retenciones": 0, "Total": 620_000,
        "CUFE_Referencia": "CUFE-FICTICIO-DEMO-005", "Estado": "Validado",
        "RutaArchivo": "DEMO://documentos/FACTURA_DEMO_005.xml", "NombreArchivo": "FACTURA_DEMO_005.xml",
        "TipoArchivo": "XML", "FechaExtraccion": datetime(2026, 9, 21, 8, 45), "MetodoExtraccion": "XML",
        "Confianza": 1, "Observaciones": "Factura ficticia relacionada con cuenta pendiente de pago.",
    },
    {
        "ID": "DOC-DEMO-006", "Fecha": date(2026, 9, 23), "TipoDocumento": "Recibo", "Numero": "RC-DEMO-006",
        "Tercero": "Nube Clara Demo S.A.S.", "Identificacion": "900000003-4 (FICTICIO)",
        "Subtotal": 185_000, "Impuestos": 0, "Retenciones": 0, "Total": 185_000,
        "CUFE_Referencia": "", "Estado": "Revision Manual",
        "RutaArchivo": "DEMO://documentos/RECIBO_DEMO_006.jpg", "NombreArchivo": "RECIBO_DEMO_006.jpg",
        "TipoArchivo": "JPG", "FechaExtraccion": datetime(2026, 9, 23, 16, 20), "MetodoExtraccion": "OCR simulado",
        "Confianza": 0.64, "Observaciones": "Confianza baja en fecha y tercero: revisar el soporte antes de registrarlo.",
    },
]

PAYMENTS = [
    {"ID": "PAG-DEMO-001", "Fecha": date(2026, 9, 10), "MovimientoID": "MOV-DEMO-003", "Tercero": "Suministros Pacífico Demo S.A.S.", "Cuenta": "Banco Principal", "MedioPago": "Transferencia", "ValorPagado": 850_000, "Estado": "Aplicado", "Observaciones": "Pago ficticio aplicado al movimiento de compra."},
    {"ID": "PAG-DEMO-002", "Fecha": date(2026, 9, 18), "MovimientoID": "MOV-DEMO-004", "Tercero": "Suministros Pacífico Demo S.A.S.", "Cuenta": "Banco Principal", "MedioPago": "Transferencia", "ValorPagado": 420_000, "Estado": "Aplicado", "Observaciones": "Pago ficticio aplicado al movimiento de compra."},
]

COLLECTIONS = [
    {"ID": "COB-DEMO-001", "Fecha": date(2026, 9, 8), "MovimientoID": "MOV-DEMO-001", "Tercero": "Cliente Aurora Demo S.A.S.", "Cuenta": "Banco Principal", "MedioPago": "Transferencia", "ValorCobrado": 4_800_000, "Estado": "Aplicado", "Observaciones": "Cobro ficticio aplicado al movimiento de ingreso."},
]


def headers_and_first_row(ws, table_name: str) -> tuple[list[str], int]:
    table = ws.tables[table_name]
    min_col, header_row, max_col, _ = range_boundaries(table.ref)
    headers = [str(ws.cell(header_row, col).value) for col in range(min_col, max_col + 1)]
    return headers, header_row + 1


def set_table_rows(wb, sheet: str, table_name: str, rows: list[dict]) -> None:
    ws = wb[sheet]
    headers, first_row = headers_and_first_row(ws, table_name)
    for offset, row_data in enumerate(rows):
        row_number = first_row + offset
        for key, value in row_data.items():
            if key in headers:
                ws.cell(row_number, headers.index(key) + 1, value)


def set_company(wb) -> None:
    ws = wb["EMPRESA"]
    headers, first_row = headers_and_first_row(ws, "EMPRESA")
    key_col = headers.index("Campo") + 1
    value_col = headers.index("Valor") + 1
    for row in range(first_row, ws.max_row + 1):
        key = ws.cell(row, key_col).value
        if key in COMPANY:
            ws.cell(row, value_col, COMPANY[key])


def set_pending_report_formulas(wb) -> None:
    """Keep receivables/payables useful in Excel 2019 using AGGREGATE formulas."""
    criteria = {
        "CXC": '(MOVIMIENTOS!$C$5:$C$10004="Ingreso")+(MOVIMIENTOS!$C$5:$C$10004="Venta")',
        "CXP": '(MOVIMIENTOS!$C$5:$C$10004="Gasto")+(MOVIMIENTOS!$C$5:$C$10004="Compra")',
    }
    columns = {"A": "$A", "B": "$B", "C": "$E", "D": "$D", "E": "$O", "G": "$P"}
    for sheet_name, type_test in criteria.items():
        ws = wb[sheet_name]
        if "A5:G5" in {str(merged) for merged in ws.merged_cells.ranges}:
            ws.unmerge_cells("A5:G5")
        ws["A5"] = None
        for row in range(5, 105):
            nth = f"ROWS($A$5:A{row})"
            pos = (
                f'AGGREGATE(15,6,(ROW(MOVIMIENTOS!$P$5:$P$10004)-ROW(MOVIMIENTOS!$P$5)+1)'
                f'/(({type_test})*(MOVIMIENTOS!$P$5:$P$10004="Pendiente")),{nth})'
            )
            for col, source_col in columns.items():
                ws[f"{col}{row}"] = f'=IFERROR(INDEX(MOVIMIENTOS!{source_col}$5:{source_col}$10004,{pos}),"")'
            ws[f"F{row}"] = f'=IF(A{row}="","",MAX(0,TODAY()-B{row}))'


def add_audit_log(wb) -> None:
    events = [
        ("Configuración ficticia de empresa", "EMPRESA", "Comercial Andina Demo S.A.S.", "LISTO"),
        ("Catálogos revisados", "CONFIGURACIÓN", "Categorías e impuestos configurables", "LISTO"),
        ("Documento importado", "DOCUMENTOS", "DOC-DEMO-001 · XML", "VALIDADO"),
        ("Documento importado", "DOCUMENTOS", "DOC-DEMO-006 · imagen", "REVISIÓN MANUAL"),
        ("Terceros identificados", "TERCEROS", "3 registros ficticios", "LISTO"),
        ("Movimientos preparados", "MOVIMIENTOS", "5 registros ficticios", "LISTO"),
        ("Pagos y cobros aplicados", "TESORERÍA", "2 pagos · 1 cobro", "LISTO"),
        ("Libro de demostración generado", "REPORTES", OUTPUT.name, "COMPLETO"),
    ]
    ws = wb["LOG"]
    headers, first_row = headers_and_first_row(ws, "LOG")
    now = datetime(2026, 10, 1, 12, 0, 0)
    for i, (action, module, affected, result) in enumerate(events):
        row = {
            "Fecha": now.date(), "Hora": time(12, i, 0), "Usuario": "DEMO",
            "Accion": action, "Modulo": module, "RegistroAfectado": affected,
            "ValorAnterior": "", "ValorNuevo": affected, "Resultado": result,
        }
        row_number = first_row + i
        for key, value in row.items():
            if key in headers:
                ws.cell(row_number, headers.index(key) + 1, value)


def create_workbook() -> dict:
    if not SOURCE.is_file():
        raise FileNotFoundError(f"No se encontró la plantilla financiera: {SOURCE}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    wb.properties.creator = "Sistema Financiero · Entorno de demostración"
    wb.properties.title = "Simulación ficticia del Sistema Financiero Empresarial"
    wb.properties.subject = "Libro de demostración generado desde la plantilla del proyecto"
    wb["INICIO"]["A1"] = "SISTEMA FINANCIERO · DEMO — DATOS FICTICIOS"
    wb["INICIO"]["A2"] = "LIBRO DE SIMULACIÓN. No usar para registrar operaciones reales ni como liquidación tributaria."
    set_company(wb)
    set_table_rows(wb, "TERCEROS", "TERCEROS", THIRD_PARTIES)
    set_table_rows(wb, "DOCUMENTOS", "DOCUMENTOS", DOCUMENTS)
    set_table_rows(wb, "MOVIMIENTOS", "MOVIMIENTOS", MOVEMENTS)
    set_table_rows(wb, "PAGOS", "PAGOS", PAYMENTS)
    set_table_rows(wb, "COBROS", "COBROS", COLLECTIONS)

    for sheet_name in ("INGRESOS", "GASTOS", "FLUJO_CAJA", "IMPUESTOS_REPORTE", "REPORTE_MENSUAL"):
        wb[sheet_name]["B3"] = 2026
    wb["REPORTE_ANUAL"]["B3"] = 2026
    wb["DASHBOARD"]["B5"] = 2026
    wb["DASHBOARD"]["B6"] = 9
    set_pending_report_formulas(wb)
    add_audit_log(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT)

    summary = {
        "company": COMPANY["NombreEmpresa"],
        "period": "Septiembre de 2026",
        "workbook": OUTPUT.name,
        "download": "/demo/" + OUTPUT.name,
        "counts": {
            "third_parties": len(THIRD_PARTIES), "documents": len(DOCUMENTS),
            "validated_documents": sum(doc["Estado"] == "Validado" for doc in DOCUMENTS),
            "manual_review_documents": sum(doc["Estado"] == "Revision Manual" for doc in DOCUMENTS),
            "movements": len(MOVEMENTS), "payments": len(PAYMENTS), "collections": len(COLLECTIONS),
        },
        "totals": {
            "income_paid": 4_800_000, "expenses_paid": 1_270_000, "net_cash_flow": 3_530_000,
            "receivable": 2_150_000, "payable": 620_000,
        },
        "documents": DOCUMENTS,
        "movements": MOVEMENTS,
        "notice": "Todos los nombres, identificaciones, soportes e importes son ficticios. Tasas de impuestos y retenciones quedan sin calcular y se configuran por empresa.",
    }
    MANIFEST.write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return summary


if __name__ == "__main__":
    result = create_workbook()
    print(f"Libro creado: {OUTPUT}")
    print(f"Terceros: {result['counts']['third_parties']} · documentos: {result['counts']['documents']} · movimientos: {result['counts']['movements']}")
