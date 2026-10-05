import shutil
import sys
import tempfile
import unittest
import io
import uuid
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))
import server


class AdminAuditTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.path = Path(self.temp_dir.name) / "audit.xlsx"
        shutil.copy2(server.DEFAULT_WORKBOOK, self.path)
        self.admin = server.WorkbookAdmin(self.path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def seed(self, entity, values):
        config = server.ENTITY_CONFIG[entity]
        workbook = load_workbook(self.path)
        sheet = workbook[config["sheet"]]
        table = sheet.tables[config["table"]]
        min_col, header_row, _, _ = range_boundaries(table.ref)
        headers = [sheet.cell(header_row, col).value for col in range(min_col, min_col + len(table.tableColumns))]
        row = header_row + 1
        for field, value in values.items():
            sheet.cell(row, min_col + headers.index(field)).value = value
        workbook.save(self.path)
        workbook.close()
        self.admin = server.WorkbookAdmin(self.path)

    def test_ambiguous_aliases_resolve_for_each_module(self):
        self.assertEqual(self.admin._map_fields("movement", {"identificaci\u00f3n": "900.123", "n\u00famero": "FV-9", "descripci\u00f3n": "Venta"}),
                         {"NIT_ID": "900.123", "Documento": "FV-9", "Concepto": "Venta"})
        self.assertEqual(self.admin._map_fields("document", {"identificaci\u00f3n": "900.123", "tipo": "Factura", "n\u00famero": "FV-9"}),
                         {"Identificacion": "900.123", "TipoDocumento": "Factura", "Numero": "FV-9"})
        self.assertEqual(self.admin._map_fields("category", {"descripci\u00f3n": "Detalle"}), {"Descripcion": "Detalle"})
        self.assertEqual(self.admin._map_fields("third_party", {"identificacion": "900.123"}), {"NIT_ID": "900.123"})

    def test_optional_fields_can_be_cleared_explicitly(self):
        self.seed("third_party", {
            "ID": "TER-00001", "Tipo": "Proveedor", "NombreRazonSocial": "Proveedor Prueba",
            "NIT_ID": "9001234567", "Email": "contacto@example.test", "Estado": "Activo",
        })
        wb = self.admin.open()
        try:
            action = self.admin._normalize_action(wb, {
                "action": "update", "entity": "third_party", "id": "TER-00001",
                "data": {"Email": ""}, "clear_fields": ["Email"],
            })
            self.assertIsNone(action["data"]["Email"])
            with self.assertRaisesRegex(ValueError, "obligatorio"):
                self.admin._normalize_action(wb, {
                    "action": "update", "entity": "third_party", "id": "TER-00001",
                    "data": {}, "clear_fields": ["NombreRazonSocial"],
                })
        finally:
            wb.close()

    def test_nit_duplicate_ignores_punctuation(self):
        self.seed("third_party", {
            "ID": "TER-00001", "Tipo": "Proveedor", "NombreRazonSocial": "Proveedor Prueba",
            "NIT_ID": "900.123.456-7", "Estado": "Activo",
        })
        with self.assertRaisesRegex(ValueError, "Ya existe un tercero"):
            self.admin.validate_actions([{
                "action": "create", "entity": "third_party",
                "data": {"Tipo": "Proveedor", "NombreRazonSocial": "Otro nombre", "NIT_ID": "9001234567"},
            }])

    def test_name_duplicate_is_flagged_without_blocking_a_third_party(self):
        self.seed("third_party", {
            "ID": "TER-00001", "Tipo": "Proveedor", "NombreRazonSocial": "José Pérez SAS",
            "Estado": "Activo",
        })
        action = self.admin.validate_actions([{
            "action": "create", "entity": "third_party",
            "data": {"Tipo": "Proveedor", "NombreRazonSocial": "Jose Perez, S.A.S."},
        }])[0]
        self.assertIn("POSIBLE_DUPLICADO", action["data"]["Observaciones"])
        self.assertTrue(action["adjustments"])

    def test_inactive_third_party_is_not_suggested(self):
        self.seed("third_party", {
            "ID": "TER-00001", "Tipo": "Proveedor", "NombreRazonSocial": "Proveedor Inactivo",
            "Estado": "Inactivo",
        })
        self.assertNotIn("Proveedor Inactivo", self.admin.form_choices()["third_parties"])
        self.assertNotIn("Proveedor Inactivo", self.admin.ai_context()["third_parties"])

    def test_document_duplicate_is_kept_and_marked_for_review(self):
        self.seed("document", {
            "ID": "DOC-00001", "Fecha": "2026-09-01", "TipoDocumento": "Factura", "Numero": "FV-001",
            "Tercero": "Proveedor Prueba", "Identificacion": "900.123.456-7", "Subtotal": 100,
            "Impuestos": 0, "Retenciones": 0, "Total": 100, "Estado": "Pendiente",
        })
        action = self.admin.validate_actions([{
            "action": "create", "entity": "document",
            "data": {"Fecha": "2026-09-01", "TipoDocumento": "Factura", "Numero": "FV-001",
                     "Tercero": "Proveedor Prueba", "Identificacion": "9001234567", "Subtotal": 100, "Total": 100},
        }])[0]
        self.assertEqual(action["data"]["Estado"], "Duplicado")
        self.assertIn("DUPLICADO", action["data"]["Observaciones"])
        self.assertIn("se conservará", action["adjustments"][0])
        result = self.admin.apply_actions([action])
        self.assertTrue(result["results"][0]["warnings"])
        documents = self.admin.read_table("document")["rows"]
        duplicate = next(row for row in documents if row.get("ID") != "DOC-00001")
        self.assertEqual(duplicate["Estado"], "Duplicado")
        self.assertIn("DUPLICADO", duplicate["Observaciones"])

    def test_clear_field_is_written_to_a_backup_and_audit_log(self):
        self.seed("third_party", {
            "ID": "TER-00001", "Tipo": "Proveedor", "NombreRazonSocial": "Proveedor Prueba",
            "NIT_ID": "9001234567", "Email": "contacto@example.test", "Estado": "Activo",
        })
        result = self.admin.apply_actions([{
            "action": "update", "entity": "third_party", "id": "TER-00001",
            "data": {}, "clear_fields": ["Email"],
        }])
        self.assertTrue((self.path.parent / "backups" / result["backup"]).is_file())
        workbook = load_workbook(self.path, data_only=False)
        try:
            row = self.admin._existing_row(workbook, "third_party", "TER-00001")
            self.assertEqual(workbook[row[0].title].cell(row[2], row[3] + row[4].index("Email")).value, None)
            self.assertGreaterEqual(len(self.admin.audit_log()), 1)
        finally:
            workbook.close()

    def test_local_post_origin_policy(self):
        self.assertTrue(server.is_local_request("127.0.0.1:8765", "http://127.0.0.1:8765", "same-origin"))
        self.assertFalse(server.is_local_request("127.0.0.1:8765", "http://malicious.example", "cross-site"))
        self.assertFalse(server.is_local_request("localhost:8765", "http://localhost:8765", "same-origin"))

    def test_invalid_content_length_is_rejected(self):
        handler = object.__new__(server.AppHandler)
        handler.headers = {"Content-Length": "-1"}
        handler.rfile = io.BytesIO(b"")
        with self.assertRaisesRegex(ValueError, "tamaño de la solicitud"):
            handler._body()

    def test_workbook_missing_a_required_table_is_rejected(self):
        workbook = load_workbook(self.path)
        del workbook["MOVIMIENTOS"].tables["MOVIMIENTOS"]
        workbook.save(self.path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "tablas necesarias"):
            server.WorkbookAdmin(self.path)

    def test_repeated_manual_request_does_not_create_a_second_record(self):
        action = {
            "action": "create", "entity": "third_party",
            "data": {"Tipo": "Proveedor", "NombreRazonSocial": "Solicitud Idempotente"},
        }
        request_id = "test-" + uuid.uuid4().hex
        with server.WRITE_LOCK:
            first = server.apply_idempotently(self.admin, [action], request_id)
            second = server.apply_idempotently(self.admin, [action], request_id)
        self.assertEqual(first, second)
        records = self.admin.read_table("third_party")["rows"]
        self.assertEqual(sum(row.get("NombreRazonSocial") == "Solicitud Idempotente" for row in records), 1)


if __name__ == "__main__":
    unittest.main()
